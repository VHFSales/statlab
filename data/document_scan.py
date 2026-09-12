"""Scan a whole document (PDF or Word) and find every table, so a thesis/
dissertation can be dropped in and its data tables detected automatically.

Scope and honesty (aligned with StatLab's principles):
- TABLES are extracted reliably: PDF via pdfplumber (per page), Word via
  python-docx (all tables), plus a text-line fallback for PDFs without ruled
  tables.
- Each detected table is SCORED by how much it looks like a statistical data table
  (fraction of numeric cells; presence of group / mean / sd / n columns; size), so
  the UI can rank them and let the user pick.
- GRAPHS/FIGURES are NOT mined for numbers. Estimating values from the pixels of a
  boxplot or curve is unreliable and would fabricate data — the opposite of this
  project's design. We only DETECT and COUNT figures and warn the user that data
  hidden only in a chart must be obtained from the underlying table or raw values.

This module returns structured results; the UI decides what to do with them.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import List, Optional

from data.importer import (MissingReader, UnsupportedFile, parse_number, _norm,
                           _GROUP_NAMES, _MEAN_NAMES, _SD_NAMES, _N_NAMES,
                           parse_delimited_text)


@dataclass
class DetectedTable:
    """A table found in a document, with metadata and a data-likelihood score."""
    rows: List[List[str]]
    location: str                 # e.g. "PDF página 12, tabela 1" / "Word tabela 3"
    n_rows: int
    n_cols: int
    caption: str = ""             # nearby caption/title text, if any
    numeric_fraction: float = 0.0
    has_group_col: bool = False
    has_summary_cols: bool = False
    score: float = 0.0
    reasons: List[str] = field(default_factory=list)


@dataclass
class DocumentScan:
    tables: List[DetectedTable]
    n_figures: int = 0
    figure_note: str = ""
    source: str = ""


# --------------------------------------------------------------------------- #
# Scoring
# --------------------------------------------------------------------------- #
def _numeric_fraction(rows: List[List[str]]) -> float:
    """Fraction of non-empty body cells that parse as numbers."""
    body = rows[1:] if len(rows) > 1 else rows
    total = 0
    numeric = 0
    for r in body:
        for c in r:
            if str(c).strip() == "":
                continue
            total += 1
            if parse_number(c) is not None:
                numeric += 1
    return (numeric / total) if total else 0.0


def _header_flags(rows: List[List[str]]):
    if not rows:
        return False, False
    header = [_norm(h) for h in rows[0]]
    has_group = any(h in _GROUP_NAMES for h in header)
    has_mean = any(h in _MEAN_NAMES for h in header)
    has_sd = any(h in _SD_NAMES for h in header)
    has_n = any(h in _N_NAMES for h in header)
    return has_group, (has_mean and (has_sd or has_n))


def score_table(t: DetectedTable) -> DetectedTable:
    """Compute a 0..1 data-likelihood score and human-readable reasons."""
    t.numeric_fraction = _numeric_fraction(t.rows)
    t.has_group_col, t.has_summary_cols = _header_flags(t.rows)
    score = 0.0
    reasons = []
    # numeric density is the strongest signal
    score += 0.6 * t.numeric_fraction
    if t.numeric_fraction >= 0.5:
        reasons.append(f"{t.numeric_fraction*100:.0f}% das células são numéricas")
    elif t.numeric_fraction > 0:
        reasons.append(f"apenas {t.numeric_fraction*100:.0f}% das células são "
                       "numéricas")
    else:
        reasons.append("nenhuma célula numérica (provavelmente texto)")
    # a plausible statistical shape (needs enough rows AND columns)
    if t.n_rows >= 4 and t.n_cols >= 2:
        score += 0.15
        reasons.append("dimensões compatíveis com uma tabela de dados")
    elif t.n_rows >= 3 and t.n_cols >= 2:
        score += 0.07
    else:
        reasons.append("tabela pequena demais para análise")
    # recognizable columns
    if t.has_group_col:
        score += 0.1
        reasons.append("coluna de grupo/tratamento reconhecida")
    if t.has_summary_cols:
        score += 0.15
        reasons.append("colunas de média/DP/n reconhecidas")
    # hard penalty: fewer than 2 body rows cannot support any analysis
    body_rows = t.n_rows - 1
    if body_rows < 2:
        score *= 0.3
        reasons.append("poucas linhas de dados para qualquer análise")
    t.score = round(min(1.0, score), 3)
    t.reasons = reasons
    return t


# --------------------------------------------------------------------------- #
# Word (.docx)
# --------------------------------------------------------------------------- #
def _scan_docx(data: bytes) -> DocumentScan:
    try:
        import docx
    except Exception:
        raise MissingReader(
            "Ler tabelas de Word (.docx) requer o pacote python-docx. Instale com: "
            "pip install python-docx")
    document = docx.Document(io.BytesIO(data))
    tables: List[DetectedTable] = []
    for i, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                rows.append(cells)
        if not rows:
            continue
        ncols = max(len(r) for r in rows)
        dt = DetectedTable(rows=rows, location=f"Word tabela {i}",
                           n_rows=len(rows), n_cols=ncols)
        tables.append(score_table(dt))
    # count drawing elements (figures) — best-effort
    n_figs = 0
    try:
        xml = document.element.xml
        n_figs = xml.count("<w:drawing") + xml.count("<pic:pic")
    except Exception:
        n_figs = 0
    note = ""
    if n_figs:
        note = (f"{n_figs} figura(s) detectada(s). Dados presentes apenas em "
                "gráficos NÃO são extraídos (estimar valores de um gráfico é "
                "cientificamente inseguro). Use a tabela correspondente ou os dados "
                "brutos.")
    return DocumentScan(tables=tables, n_figures=n_figs, figure_note=note,
                        source="Word (.docx)")


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
_CAPTION_RE = re.compile(r"^\s*(tabela|table|quadro)\b", re.IGNORECASE)


def _scan_pdf(data: bytes) -> DocumentScan:
    try:
        import pdfplumber
    except Exception:
        raise MissingReader(
            "Ler tabelas de PDF requer o pacote pdfplumber. Instale com: "
            "pip install pdfplumber — ou exporte as tabelas como CSV/Excel.")
    tables: List[DetectedTable] = []
    n_figs = 0
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for pno, page in enumerate(pdf.pages, start=1):
            # count images/figures on the page
            try:
                n_figs += len(page.images or [])
            except Exception:
                pass
            # find a caption line on the page (best-effort)
            caption = ""
            try:
                text = page.extract_text() or ""
                for line in text.splitlines():
                    if _CAPTION_RE.match(line):
                        caption = line.strip()[:120]
                        break
            except Exception:
                text = ""
            # ruled tables
            found_any = False
            try:
                extracted = page.extract_tables() or []
            except Exception:
                extracted = []
            for ti, tbl in enumerate(extracted, start=1):
                rows = [[("" if c is None else str(c).strip()) for c in row]
                        for row in tbl if any(row)]
                if not rows:
                    continue
                ncols = max(len(r) for r in rows)
                dt = DetectedTable(
                    rows=rows,
                    location=f"PDF página {pno}, tabela {ti}",
                    n_rows=len(rows), n_cols=ncols, caption=caption)
                tables.append(score_table(dt))
                found_any = True
            # text-line fallback: only if no ruled table AND the page looks tabular
            if not found_any and text:
                rows = []
                for line in text.splitlines():
                    parts = re.split(r"\s{2,}|\t", line.strip())
                    if len(parts) >= 2:
                        rows.append([p.strip() for p in parts])
                if len(rows) >= 3:
                    dt = DetectedTable(
                        rows=rows,
                        location=f"PDF página {pno} (texto)",
                        n_rows=len(rows),
                        n_cols=max(len(r) for r in rows), caption=caption)
                    scored = score_table(dt)
                    # only keep text-fallback tables that look numeric enough
                    if scored.numeric_fraction >= 0.3:
                        tables.append(scored)
    note = ""
    if n_figs:
        note = (f"{n_figs} imagem(ns)/figura(s) detectada(s) no PDF. Dados presentes "
                "apenas em gráficos NÃO são extraídos (estimar valores de um gráfico "
                "é cientificamente inseguro). Procure a tabela correspondente.")
    return DocumentScan(tables=tables, n_figures=n_figs, figure_note=note,
                        source="PDF")


# --------------------------------------------------------------------------- #
# Public entry point
# --------------------------------------------------------------------------- #
def table_to_dataset(table: DetectedTable, kind: str = "auto",
                     decimal: str = "auto"):
    """Convert a detected table into a StatLab dataset.

    ``kind``: "raw", "summary", or "auto" (uses the header flags to decide).
    Returns (kind_used, dataset) where dataset is the raw dict or summary dict.
    """
    from data.importer import rows_to_raw, rows_to_summary
    use = kind
    if use == "auto":
        use = "summary" if table.has_summary_cols else "raw"
    if use == "summary":
        summ, _dec = rows_to_summary(table.rows, decimal)
        return "summary", summ
    raw, _fmt, _dec = rows_to_raw(table.rows, None, decimal)
    return "raw", raw


def scan_document(data: bytes, filename: str) -> DocumentScan:
    """Scan a PDF or Word document and return all detected tables (scored) plus a
    figure count. Tables are sorted by descending data-likelihood score."""
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        scan = _scan_pdf(data)
    elif name.endswith(".docx"):
        scan = _scan_docx(data)
    else:
        raise UnsupportedFile(
            "A varredura de documento aceita .pdf ou .docx. Para .xlsx/.csv, use o "
            "envio de arquivo normal na seção DADOS.")
    scan.tables.sort(key=lambda t: t.score, reverse=True)
    return scan
